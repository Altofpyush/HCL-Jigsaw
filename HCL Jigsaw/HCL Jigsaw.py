
"""
HCL Jigsaw Round 2
Traffic Flow Optimizer
Pacific World School, Greater Noida West
Team: team_965_0045
"""

from pathlib import Path
import openpyxl
import pandas as pd
import matplotlib.pyplot as plt

# --------------------------------------------------
# PROJECT SETTINGS
# --------------------------------------------------

WORKBOOK_NAME = "Traffic_Flow_Optimizer.xlsx"
OUTPUT_FOLDER = "traffic_optimizer_outputs"


# --------------------------------------------------
# MAIN PROGRAM
# --------------------------------------------------

def main():
    folder = Path(__file__).resolve().parent
    workbook_path = folder / WORKBOOK_NAME

    print("\n" + "=" * 60)
    print("HCL JIGSAW ROUND 2")
    print("TRAFFIC FLOW OPTIMIZER")
    print("Pacific World School, Greater Noida West")
    print("Team: team_965_0045")
    print("=" * 60)

    print("\nPython file folder:")
    print(folder)

    # Find the workbook if the exact name differs
    if not workbook_path.exists():
        excel_files = list(folder.glob("*.xlsx"))

        if len(excel_files) == 1:
            workbook_path = excel_files[0]
            print("\nUsing Excel workbook:", workbook_path.name)

        elif len(excel_files) == 0:
            print("\nERROR: No Excel workbook found.")
            print("Put the Excel file beside HCL Jigsaw.py.")
            print("Files in the project folder:")

            for item in folder.iterdir():
                print("-", item.name)

            return

        else:
            print("\nSeveral Excel files were found.")
            print("Please check the workbook name.")
            for item in excel_files:
                print("-", item.name)
            return

    # --------------------------------------------------
    # READ EXCEL WORKBOOK
    # --------------------------------------------------

    try:
        workbook = openpyxl.load_workbook(
            workbook_path,
            data_only=True,
            read_only=True
        )

        print("\nWorkbook opened successfully!")
        print("Workbook:", workbook_path.name)
        print("\nWorksheets found:")

        for name in workbook.sheetnames:
            print("-", name)

        output_folder = folder / OUTPUT_FOLDER
        output_folder.mkdir(exist_ok=True)

        exported_sheets = []

        # Read every worksheet without assuming sheet names
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            rows = list(sheet.iter_rows(values_only=True))

            print("\n" + "-" * 60)
            print("WORKSHEET:", sheet_name)
            print("-" * 60)

            if not rows:
                print("This worksheet is empty.")
                continue

            # Display the first few rows
            for row in rows[:8]:
                print(row)

            # Convert the worksheet into a DataFrame
            headers = list(rows[0])

            # Replace blank column headings
            for index, header in enumerate(headers):
                if header is None:
                    headers[index] = f"Column_{index + 1}"

            data = rows[1:]

            # Make duplicate column headings unique
            unique_headers = []
            for index, header in enumerate(headers):
                header = str(header)
                candidate = header
                count = 2

                while candidate in unique_headers:
                    candidate = f"{header}_{count}"
                    count += 1

                unique_headers.append(candidate)

            dataframe = pd.DataFrame(
                data,
                columns=unique_headers
            ).dropna(how="all")

            safe_name = "".join(
                character if character.isalnum()
                or character in "_-" else "_"
                for character in sheet_name
            )

            csv_path = output_folder / f"{safe_name}.csv"

            dataframe.to_csv(csv_path, index=False)
            exported_sheets.append((sheet_name, dataframe))

            print("\nExported:", csv_path.name)

        workbook.close()

    except PermissionError:
        print("\nERROR: Windows cannot access the workbook.")
        print("Close the workbook in Excel and try again.")
        return

    except Exception as error:
        print("\nERROR while reading Excel:")
        print(type(error).__name__, "-", error)
        return

    # --------------------------------------------------
    # CREATE CHARTS FROM NUMERIC WORKSHEET DATA
    # --------------------------------------------------

    charts_folder = folder / OUTPUT_FOLDER
    charts_created = 0

    for sheet_name, dataframe in exported_sheets:
        if dataframe.empty:
            continue

        numeric_data = dataframe.select_dtypes(
            include="number"
        )

        if numeric_data.empty:
            continue

        # Plot numeric columns against row number.
        # This is a data overview, not a recalculated simulation.
        plt.figure(figsize=(10, 5))

        for column in numeric_data.columns[:6]:
            plt.plot(
                range(1, len(numeric_data) + 1),
                numeric_data[column],
                label=str(column)
            )

        plt.title(f"Data Overview: {sheet_name}")
        plt.xlabel("Row Number")
        plt.ylabel("Recorded Value")
        plt.grid(True, alpha=0.3)

        if len(numeric_data.columns) <= 6:
            plt.legend()

        plt.tight_layout()

        safe_name = "".join(
            character if character.isalnum()
            or character in "_-" else "_"
            for character in sheet_name
        )

        chart_path = charts_folder / f"{safe_name}_chart.png"

        plt.savefig(chart_path, dpi=160)
        plt.close()

        print("\nChart created:", chart_path.name)
        charts_created += 1

    # --------------------------------------------------
    # FINISH
    # --------------------------------------------------

    print("\n" + "=" * 60)
    print("PROGRAM FINISHED")
    print("=" * 60)

    print("Excel workbook:", workbook_path.name)
    print("Worksheets exported:", len(exported_sheets))
    print("Charts created:", charts_created)
    print("Output folder:", charts_folder)

    print("\nYour CSV files and charts are saved here:")
    print(charts_folder.resolve())


if __name__ == "__main__":
    main()
